"""Import real Horti-M3 tables without modifying the archive or source workbooks."""
import argparse
import csv
import hashlib
import io
import json
import math
from collections import Counter
from datetime import datetime, timedelta, timezone
from pathlib import Path

import openpyxl

VERSION = 'horti-m3-2025-observations-v1'
START, END = '2025-04-19', '2025-06-13'
PLANTS = ['CK11', 'CK12', 'CK13', 'CK14', 'CK15', 'CK16']
PAPER = 'https://www.nature.com/articles/s41597-026-07074-w'
METRICS = {
    '株高': ('plantHeightCm', 'cm', 'measured'),
    '茎粗': ('stemDiameterMm', 'mm', 'measured'),
    'LAI': ('canopyLai', 'm2/m2; canopy instrument basis', 'instrument_estimated'),
    'LDW': ('canopyLdw', 'g; canopy instrument estimate', 'instrument_estimated'),
}
SENSOR_COLUMNS = {
    'temperatureC': '空气温度（对照2）后',
    'airHumidityPct': '相对湿度（对照2）后',
    'co2Ppm': '二氧化碳（对照2）后',
    'lightRaw': '光照（对照2）后',
    'soilMoistureVwcPct': '土壤水分（对照2）',
    'soilTemperatureC': '土壤温度（对照2）',
}


def number(value):
    try:
        n = float(value)
        return n if math.isfinite(n) else None
    except (ValueError, TypeError):
        return None


def write_json(path, data):
    temporary = path.with_suffix('.tmp')
    temporary.write_text(json.dumps(data, ensure_ascii=False, allow_nan=False, indent=2), encoding='utf-8')
    temporary.replace(path)


def import_tables(dataset_dir):
    root = dataset_dir / 'tables' / '2023-2025 Tomato dataset'
    if not root.is_dir():
        raise ValueError('Extract the real data tables before importing.')
    status = json.loads((dataset_dir / 'download-status.json').read_text(encoding='utf-8-sig'))
    if status.get('checksumVerified') is not True or status.get('stage') != 'complete':
        raise ValueError('Official archive checksum must be verified before importing.')
    output = dataset_dir / 'normalized'
    output.mkdir(exist_ok=True)
    sources = []
    for p in sorted(root.rglob('*')):
        if p.is_file():
            relative = p.relative_to(root).as_posix()
            sources.append({'file': relative, 'bytes': p.stat().st_size,
                            'sha256': hashlib.sha256(p.read_bytes()).hexdigest(),
                            'kind': 'provider_interpolated' if 'after interpolation' in relative
                            else 'provider_combined' if 'growth_index(final)' in relative
                            else 'original_source'})

    growth, issues = [], []
    for p in sorted(root.glob('2025/**/raw_data/growth_index/*.csv')):
        text = p.read_bytes().decode('gb18030')
        day = datetime.strptime(p.stem, '%Y%m%d').date().isoformat()
        reader = csv.DictReader(io.StringIO(text))
        for row in reader:
            plant = row.get('编号', '').strip()
            if not plant:
                continue
            for column, (metric, unit, kind) in METRICS.items():
                value = number(row.get(column))
                flags = []
                if value is None:
                    flags.append('MISSING_OR_NONFINITE')
                elif value < 0 or (metric == 'plantHeightCm' and value > 600):
                    flags.append('OUT_OF_RANGE')
                growth.append({'observedDate': day, 'timePrecision': 'date', 'plantId': plant,
                               'metric': metric, 'value': value, 'unit': unit, 'observationKind': kind,
                               'selected': plant in PLANTS, 'qualityFlags': flags,
                               'source': {'file': p.relative_to(root).as_posix(), 'row': reader.line_num,
                                          'column': column, 'encoding': 'gb18030'}, 'raw': dict(row)})

    environment, raw_environment = [], []
    keyed = {}
    for p in sorted(root.glob('2025/**/raw_data/sensor_data/sensor1.xlsx')):
        book = openpyxl.load_workbook(p, read_only=True, data_only=True)
        for sheet in book:
            rows = sheet.iter_rows(values_only=True)
            header = list(next(rows))
            indices = {k: header.index(v) for k, v in SENSOR_COLUMNS.items()}
            for line, row in enumerate(rows, start=2):
                stamp = row[0]
                if not isinstance(stamp, datetime) or not START <= stamp.date().isoformat() <= END:
                    continue
                values = {k: number(row[i]) for k, i in indices.items()}
                flags = []
                ranges = {'temperatureC': (-40, 60), 'airHumidityPct': (0, 100),
                          'co2Ppm': (150, 5000), 'soilMoistureVwcPct': (0, 100), 'lightRaw': (0, 200000)}
                for k, (low, high) in ranges.items():
                    if values[k] is None or not low <= values[k] <= high:
                        flags.append('INVALID_' + k)
                # Zero CO2 indicates a failed measurement, rather than a zero-carbon atmosphere.
                record = {'observedAt': stamp.isoformat() + '+08:00', **values,
                          'sensorId': 'sensor1-contrast2-back', 'observationKind': 'measured',
                          'qualityFlags': flags, 'source': {'file': p.relative_to(root).as_posix(),
                          'sheet': sheet.title, 'row': line, 'columns': SENSOR_COLUMNS},
                          'raw': {str(h): (v.isoformat() if isinstance(v, datetime) else v)
                                  for h, v in zip(header, row)}}
                raw_environment.append(record)
                key = record['observedAt']
                if key in keyed:
                    old = keyed[key]
                    same = all(old[k] == values[k] for k in indices)
                    issues.append({'kind': 'EXACT_DUPLICATE' if same else 'CONFLICTING_DUPLICATE',
                                   'at': key, 'source': record['source']})
                    if not same:
                        old['qualityFlags'].append('CONFLICTING_DUPLICATE')
                else:
                    keyed[key] = record
        book.close()
    environment = sorted(keyed.values(), key=lambda r: r['observedAt'])
    references = []
    for sensor, label in [('sensor2', '（对照2）'), ('sensor3', '（处理1）后')]:
        columns = {'temperatureC': '空气温度' + label, 'airHumidityPct': '相对湿度' + label,
                   'co2Ppm': '二氧化碳' + label, 'lightRaw': '光照' + label}
        unique = {}
        for p in sorted(root.glob(f'2025/**/raw_data/sensor_data/{sensor}.xlsx')):
            book = openpyxl.load_workbook(p, read_only=True, data_only=True)
            for sheet in book:
                rows = sheet.iter_rows(values_only=True)
                header = list(next(rows))
                if not all(c in header for c in columns.values()):
                    continue
                indices = {k: header.index(c) for k, c in columns.items()}
                for line, row in enumerate(rows, start=2):
                    stamp = row[0]
                    if not isinstance(stamp, datetime) or not START <= stamp.date().isoformat() <= END:
                        continue
                    values = {k: number(row[i]) for k, i in indices.items()}
                    if (any(v is None for v in values.values()) or not -40 <= values['temperatureC'] <= 60
                            or not 0 <= values['airHumidityPct'] <= 100 or not 150 <= values['co2Ppm'] <= 5000
                            or not 0 <= values['lightRaw'] <= 200000):
                        continue
                    key = stamp.isoformat() + '+08:00'
                    record = {'observedAt': key, **values, 'sensorId': sensor + label,
                              'source': {'file': p.relative_to(root).as_posix(), 'sheet': sheet.title,
                                         'row': line, 'columns': columns},
                              'referenceOnly': True, 'soilMoistureVwcPct': None}
                    if key in unique and any(unique[key][k] != values[k] for k in values):
                        unique[key]['conflict'] = True
                    else:
                        unique.setdefault(key, record)
            book.close()
        references.extend(r for r in unique.values() if not r.get('conflict'))
    valid = [r for r in environment if not r['qualityFlags']]
    coverage = Counter(r['observedAt'][:10] for r in valid)
    day = datetime.fromisoformat(START)
    gaps = []
    while day.date().isoformat() <= END:
        date = day.date().isoformat()
        gaps.append({'date': date, 'validSlots': len({r['observedAt'][11:16] for r in valid
                      if r['observedAt'].startswith(date)}), 'expectedSlots': 48,
                     'validRecords': coverage[date]})
        day += timedelta(days=1)
    warnings = [
        'CK11–CK16品种依据README编号规则；小区重复与sensor1对照2后无法逐株确认，采用CK区域代表环境。',
        '株高cm、茎粗mm、LDW仪器估计g、土壤水分VWC%依据论文表9/10；LAI冠层测量分母未映射为模型地面积。',
        '光照原列无单位，论文表9写lx但表3量程/分辨率为klux；原始白天数值多为个位至几十。保留原值，运行中须明确选择分析假设。',
        '2025生长记录Sheet1含2023日期；Sheet2与农事记录部分物候日期冲突，本轮不拟合物候。',
        '原始表型仅日期精度；对齐12:00为项目约定，不是声称测量在12:00进行。',
        '不以每日插值、最终整合CSV或重复记录增加独立测量样本。',
    ]
    quality = {'rawEnvironmentRows': len(raw_environment), 'uniqueEnvironmentRows': len(environment),
               'validEnvironmentRows': len(valid), 'invalidEnvironmentRows': len(environment)-len(valid),
               'duplicateRows': len(issues), 'duplicateDetails': issues,
               'dailyCoverage': gaps, 'coveragePct': round(sum(g['validSlots'] for g in gaps)/(56*48)*100, 3),
               'warnings': warnings}
    payload = {'schemaVersion': 1, 'datasetVersion': VERSION,
               'importedAt': datetime.now(timezone(timedelta(hours=8))).isoformat(),
               'sourceUrl': 'https://zenodo.org/records/17217565', 'paperUrl': PAPER, 'license': 'CC BY 4.0',
               'publisherMd5': status['publisherMd5'], 'archiveBytes': status['expectedBytes'],
               'cohort': {'id': '2025-CK-Guanghui201', 'year': 2025, 'treatment': 'CK',
                          'variety': '广辉201', 'plantIds': PLANTS, 'startDate': START, 'endDate': END,
                          'environmentScope': 'CK区域代表环境；未确认逐株/重复小区对应'},
               'growth': growth, 'environment': environment, 'environmentReferences': references,
               'sources': sources, 'quality': quality,
               'fieldMappings': {'growth': METRICS, 'environment': SENSOR_COLUMNS}}
    write_json(output / 'environment-original.json', raw_environment)
    write_json(output / 'quality-report.json', quality)
    write_json(output / 'observations.json', payload)
    sha = hashlib.sha256((output/'observations.json').read_bytes()).hexdigest()
    write_json(output/'manifest.json', {'datasetVersion': VERSION, 'observationsSha256': sha,
               'sourceFiles': len(sources), 'archiveMd5': status['publisherMd5']})
    print(json.dumps({'sources': len(sources), 'growthRecords': len(growth),
                      'selectedHeightObservations': sum(r['selected'] and r['metric']=='plantHeightCm' for r in growth),
                      'validEnvironmentRecords': len(valid), 'coveragePct': quality['coveragePct'],
                      'observationsSha256': sha}, ensure_ascii=False))


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('--dataset-dir', type=Path, required=True)
    import_tables(parser.parse_args().dataset_dir.resolve())
